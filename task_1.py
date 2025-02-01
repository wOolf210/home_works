import json
import os
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

# Работа с Excel
files = ['file1.xlsx', 'file2.xlsx', 'file3.xlsx']
data = []

for file in files:
    with open(file, 'w') as f:
        value = int(file[4]) * 1111
        f.write(str(value))
        data.append(value)

data.sort(reverse=True)

wb = Workbook()
ws = wb.active

for r_idx, value in enumerate(data, start=1):
    cell = ws.cell(row=r_idx, column=1, value=value)
    cell.font = Font(name='Arial', size=12, bold=True)
    cell.border = Border(left=Side(style='thin'), right=Side(style='thin'),
                         top=Side(style='thin'), bottom=Side(style='thin'))
wb.save("result.xlsx")

# Работа с JSON
json_url = "https://jsonplaceholder.typicode.com/todos/"
response = requests.get(json_url)
data = response.json()

with open("todos.json", "w", encoding="utf-8") as f:
    json.dump(data, f, indent=4, ensure_ascii=False)

with open("todos.json", "r", encoding="utf-8") as f:
    todos = json.load(f)

os.makedirs("todos", exist_ok=True)

for i, todo in enumerate(todos):
    with open(f"todos/todo_{i + 1}.json", "w", encoding="utf-8") as f:
        json.dump(todo, f, indent=4, ensure_ascii=False)

# Работа с Word
word_file = "hello.docx"
doc = Document()
doc.add_paragraph("Hello Python").runs[0].bold = True
doc.save(word_file)

new_doc = Document(word_file)
bold_text = " ".join(run.text for para in new_doc.paragraphs for run in para.runs if run.bold)
print(bold_text)

new_doc = Document()
para = new_doc.add_paragraph("This is a new paragraph with changed font size.")
run = para.runs[0]
run.font.name = "Times New Roman"
run.font.size = Pt(16)
para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
new_doc.save("new_word.docx")
